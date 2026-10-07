"""
core/exporter.py
----------------
Menghasilkan laporan rekonsiliasi Excel multi-tab profesional siap audit
dan CSV draft jurnal penyesuaian (Mekari Jurnal / Accurate Online).
"""

import io
import csv
from typing import Dict, Any, List
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

def generate_reconciliation_excel(result: Dict[str, Any]) -> io.BytesIO:
    """
    Membuat file Excel .xlsx multi-tab lengkap dengan formatting finansial profesional.
    """
    wb = openpyxl.Workbook()
    # Hapus sheet default
    wb.remove(wb.active)
    
    summary = result.get("summary", {})
    batches = result.get("batches", [])
    discrepancies = result.get("discrepancies", [])
    bank_rows = result.get("bank_rows", [])

    # Styles
    font_header = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
    font_bold = Font(name="Segoe UI", size=11, bold=True)
    font_regular = Font(name="Segoe UI", size=10)
    font_title = Font(name="Segoe UI", size=14, bold=True, color="1E3A8A")
    
    fill_navy = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
    fill_blue = PatternFill(start_color="2563EB", end_color="2563EB", fill_type="solid")
    fill_green = PatternFill(start_color="DCFCE7", end_color="DCFCE7", fill_type="solid")
    fill_red = PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid")
    fill_yellow = PatternFill(start_color="FEF9C3", end_color="FEF9C3", fill_type="solid")
    fill_gray = PatternFill(start_color="F3F4F6", end_color="F3F4F6", fill_type="solid")
    
    border_thin = Border(
        left=Side(style='thin', color='D1D5DB'),
        right=Side(style='thin', color='D1D5DB'),
        top=Side(style='thin', color='D1D5DB'),
        bottom=Side(style='thin', color='D1D5DB')
    )
    align_center = Alignment(horizontal='center', vertical='center')
    align_right = Alignment(horizontal='right', vertical='center')
    align_left = Alignment(horizontal='left', vertical='center')

    # =========================================================
    # TAB 1: EXECUTIVE SUMMARY & PROOF OF CASH
    # =========================================================
    ws1 = wb.create_sheet(title="Ringkasan & Proof of Cash")
    ws1.views.sheetView[0].showGridLines = True
    
    ws1.cell(row=2, column=2, value="LAPORAN REKONSILIASI DUA ARAH MULTI-SUMBER").font = font_title
    ws1.cell(row=3, column=2, value="Marketplace, Payment Gateway, Mutasi Bank, & Kasir Tunai").font = font_regular
    
    # Proof of Cash Box
    proof_rows = [
        ("Komponen Rekonsiliasi", "Nominal (IDR)", "Keterangan"),
        ("1. Total Penjualan Kotor (Gross Sales)", summary.get("total_gross_sales", 0), "Marketplace + DOKU + Kasir"),
        ("2. Potongan Komisi & Fee Platform", -summary.get("total_platform_fees", 0), "Referral, Transaksi, MDR"),
        ("3. Estimasi Penjualan Bersih (Net Expected)", summary.get("total_expected_net", 0), "Dana yang berhak dicairkan"),
        ("(-) Dana Dalam Perjalanan (Deposit In-Transit)", -summary.get("in_transit_amount", 0), "Kliring H+1 sd H+3"),
        ("(-) Dana Tertahan / Hilang (Missing Payout)", -summary.get("missing_payout_amount", 0), "Perlu klaim ke Seller Center"),
        ("(-) Kas Tunai Belum Disetor (Cash in Drawer)", -summary.get("unsettled_cash_amount", 0), "Fisik masih di brankas toko"),
        ("(-) Selisih Kurang Setor Kasir (Cash Shortage)", -summary.get("cash_shortage", 0), "Uang kasir kurang"),
        ("(+) Setoran Bank Tak Dikenal (Unidentified Credits)", summary.get("unidentified_bank_credits", 0), "Uang masuk tanpa pesanan"),
        ("(+) Pendapatan Bunga Bank", summary.get("bank_interest_income", 0), "Bunga rekening giro"),
        ("= Nilai Rekonsiliasi Kredit Mutasi Bank", summary.get("reconciled_bank_credits", 0), "Hasil perhitungan rumus"),
        ("Riil Total Kredit di Rekening Bank BCA", summary.get("bank_total_credits", 0), "Total mutasi CR rekening"),
        ("STATUS BUKTI REKONSILIASI (SELISIH)", summary.get("proof_difference", 0), "HARUS Rp 0,00 (TIE / SEIMBANG)")
    ]
    
    start_row = 5
    for i, (title, amt, note) in enumerate(proof_rows):
        r = start_row + i
        c1 = ws1.cell(row=r, column=2, value=title)
        c2 = ws1.cell(row=r, column=3, value=amt if isinstance(amt, str) else round(amt, 2))
        c3 = ws1.cell(row=r, column=4, value=note)
        
        for c in [c1, c2, c3]:
            c.border = border_thin
            c.font = font_regular
            
        if i == 0:
            for c in [c1, c2, c3]:
                c.fill = fill_navy
                c.font = font_header
                c.alignment = align_center
        elif i == len(proof_rows) - 1:
            for c in [c1, c2, c3]:
                c.fill = fill_green if summary.get("is_balanced") else fill_red
                c.font = font_bold
            c2.number_format = '#,##0.00'
        else:
            if isinstance(amt, (int, float)):
                c2.number_format = '#,##0.00'
                c2.alignment = align_right

    # =========================================================
    # TAB 1.5: ANALISIS MATCH VS NOT MATCH (%)
    # =========================================================
    match_analytics = result.get("match_analytics", {})
    overall = match_analytics.get("overall", {})
    by_channel = match_analytics.get("by_channel", [])
    categories = match_analytics.get("categories", [])

    ws_an = wb.create_sheet(title="Analisis Match vs Not Match")
    ws_an.views.sheetView[0].showGridLines = True

    ws_an.cell(row=2, column=2, value="ANALISIS DETAIL PERSENTASE MATCH VS NOT MATCH").font = font_title
    ws_an.cell(row=3, column=2, value="Ringkasan dan Rincian Tingkat Keberhasilan Rekonsiliasi Finansial").font = font_regular

    # Sub-tabel 1: Ringkasan Agregat
    ws_an.cell(row=5, column=2, value="1. RINGKASAN AGREGAT (OVERALL MATCH VS NOT MATCH)").font = font_bold
    headers_ov = ["Dimensi Metrik", "Matched (Klop)", "Not Matched (Selisih)", "Total Basis", "Persentase Match (%)", "Persentase Not Match (%)"]
    ws_an.append([])
    ws_an.append([""] + headers_ov)
    h_row = ws_an.max_row
    for c_idx in range(2, 2 + len(headers_ov)):
        cell = ws_an.cell(row=h_row, column=c_idx)
        cell.font = font_header
        cell.fill = fill_navy
        cell.alignment = align_center
        cell.border = border_thin

    row_nom = ["Nilai Uang (Nominal IDR)", overall.get("matched_nominal", 0), overall.get("not_matched_nominal", 0), overall.get("total_nominal", 0), f"{overall.get('matched_nominal_pct', 0)}%", f"{overall.get('not_matched_nominal_pct', 0)}%"]
    row_cnt = ["Jumlah Transaksi (Count)", overall.get("matched_count", 0), overall.get("not_matched_count", 0), overall.get("total_count", 0), f"{overall.get('matched_count_pct', 0)}%", f"{overall.get('not_matched_count_pct', 0)}%"]
    
    for r_vals in [row_nom, row_cnt]:
        ws_an.append([""] + r_vals)
        cur = ws_an.max_row
        for c_idx in range(2, 2 + len(headers_ov)):
            c = ws_an.cell(row=cur, column=c_idx)
            c.border = border_thin
            c.font = font_regular
            if c_idx in [3, 4, 5] and "Nominal" in r_vals[0]:
                c.number_format = '#,##0.00'
                c.alignment = align_right
            elif c_idx in [6, 7]:
                c.alignment = align_center
                c.font = font_bold

    # Sub-tabel 2: Per Kanal
    ws_an.append([])
    start_ch = ws_an.max_row + 1
    ws_an.cell(row=start_ch, column=2, value="2. PERSENTASE MATCH PER KANAL SUMBER (CHANNEL BREAKDOWN)").font = font_bold
    headers_ch = ["Kanal / Sumber", "Total Item", "Matched Item", "Not Matched Item", "% Match (Count)", "Total Nominal (Rp)", "Matched Nominal (Rp)", "Not Matched (Rp)", "% Match (Nominal)"]
    ws_an.append([""] + headers_ch)
    h_row = ws_an.max_row
    for c_idx in range(2, 2 + len(headers_ch)):
        cell = ws_an.cell(row=h_row, column=c_idx)
        cell.font = font_header
        cell.fill = fill_blue
        cell.alignment = align_center
        cell.border = border_thin

    for ch in by_channel:
        r_ch = [
            ch.get("channel_name"), ch.get("total_count"), ch.get("matched_count"), ch.get("not_matched_count"),
            f"{ch.get('match_count_pct', 0)}%", ch.get("total_nominal"), ch.get("matched_nominal"), ch.get("not_matched_nominal"),
            f"{ch.get('match_nominal_pct', 0)}%"
        ]
        ws_an.append([""] + r_ch)
        cur = ws_an.max_row
        for c_idx in range(2, 2 + len(headers_ch)):
            c = ws_an.cell(row=cur, column=c_idx)
            c.border = border_thin
            c.font = font_regular
            if c_idx in [7, 8, 9]:
                c.number_format = '#,##0.00'
                c.alignment = align_right
            elif c_idx in [3, 4, 5, 6, 10]:
                c.alignment = align_center
                if c_idx == 10:
                    c.font = font_bold

    # Sub-tabel 3: Rincian Kategori Status
    ws_an.append([])
    start_cat = ws_an.max_row + 1
    ws_an.cell(row=start_cat, column=2, value="3. MATRIKS RINCI KATEGORI STATUS MATCH & NOT MATCH").font = font_bold
    headers_cat = ["Kategori Status", "Tipe Status", "Jumlah Item", "% Jumlah", "Total Nominal (Rp)", "% Nominal", "Rekomendasi Tindakan"]
    ws_an.append([""] + headers_cat)
    h_row = ws_an.max_row
    for c_idx in range(2, 2 + len(headers_cat)):
        cell = ws_an.cell(row=h_row, column=c_idx)
        cell.font = font_header
        cell.fill = fill_navy
        cell.alignment = align_center
        cell.border = border_thin

    for cat in categories:
        r_cat = [
            cat.get("category"), cat.get("type"), cat.get("count"), f"{cat.get('count_pct', 0)}%",
            cat.get("nominal"), f"{cat.get('nominal_pct', 0)}%", cat.get("action")
        ]
        ws_an.append([""] + r_cat)
        cur = ws_an.max_row
        for c_idx in range(2, 2 + len(headers_cat)):
            c = ws_an.cell(row=cur, column=c_idx)
            c.border = border_thin
            c.font = font_regular
            if c_idx == 6:
                c.number_format = '#,##0.00'
                c.alignment = align_right
            elif c_idx in [3, 4, 5, 7]:
                c.alignment = align_center
                if c_idx == 3:
                    c.fill = fill_green if cat.get("type") == "MATCHED" else fill_yellow

    # =========================================================
    # TAB 2: BATCH PAYOUT TO BANK (FORWARD MATCHING)
    # =========================================================
    ws2 = wb.create_sheet(title="Forward Batches to Bank")
    ws2.views.sheetView[0].showGridLines = True
    
    headers2 = ["Batch / Ref ID", "Channel", "Tanggal Settle", "Qty Order", "Gross (Rp)", "Fee (Rp)", "Net Settle (Rp)", "Status", "Tgl Bank", "Nominal Bank (Rp)", "Selisih", "Catatan"]
    ws2.append([])
    ws2.append([""] + headers2)
    
    for c_idx in range(2, 2 + len(headers2)):
        cell = ws2.cell(row=2, column=c_idx)
        cell.font = font_header
        cell.fill = fill_blue
        cell.alignment = align_center
        cell.border = border_thin
        
    for b in batches:
        r_vals = [
            b.get("batch_id"), b.get("channel"), b.get("settlement_date"), b.get("order_count"),
            b.get("gross_sum"), b.get("fee_sum"), b.get("net_sum"), b.get("status"),
            b.get("bank_date") or "-", b.get("bank_amount") or 0.0, b.get("difference"), b.get("notes")
        ]
        ws2.append([""] + r_vals)
        cur_row = ws2.max_row
        
        for c_idx in range(2, 2 + len(headers2)):
            cell = ws2.cell(row=cur_row, column=c_idx)
            cell.font = font_regular
            cell.border = border_thin
            # Format angka
            if c_idx in [6, 7, 8, 11, 12]:
                cell.number_format = '#,##0.00'
                cell.alignment = align_right
            elif c_idx in [2, 3, 4, 5, 9, 10]:
                cell.alignment = align_center

    # =========================================================
    # TAB 3: REVERSE BANK AUDIT
    # =========================================================
    ws3 = wb.create_sheet(title="Reverse Bank Audit")
    ws3.views.sheetView[0].showGridLines = True
    
    headers3 = ["ID Baris", "Tanggal", "Keterangan Mutasi Bank", "Jumlah (IDR)", "Tipe", "Matched?", "Pasangan Batch Ref"]
    ws3.append([])
    ws3.append([""] + headers3)
    
    for c_idx in range(2, 2 + len(headers3)):
        cell = ws3.cell(row=2, column=c_idx)
        cell.font = font_header
        cell.fill = fill_navy
        cell.alignment = align_center
        cell.border = border_thin
        
    for row in bank_rows:
        r_vals = [
            row.get("row_id"), row.get("date"), row.get("description"),
            row.get("amount"), row.get("txn_type"),
            "YA (Klop)" if row.get("matched") else "TIDAK (Anomali / Perlu Jurnal)",
            row.get("matched_with") or "-"
        ]
        ws3.append([""] + r_vals)
        cur_row = ws3.max_row
        for c_idx in range(2, 2 + len(headers3)):
            cell = ws3.cell(row=cur_row, column=c_idx)
            cell.font = font_regular
            cell.border = border_thin
            if c_idx == 5:
                cell.number_format = '#,##0.00'
                cell.alignment = align_right
            elif c_idx == 7:
                cell.fill = fill_green if row.get("matched") else fill_yellow

    # =========================================================
    # TAB 4: EXCEPTIONS & ACTION ITEMS
    # =========================================================
    ws4 = wb.create_sheet(title="Exceptions & Tindakan")
    ws4.views.sheetView[0].showGridLines = True
    
    headers4 = ["Arah Uji", "Channel", "Jenis Masalah", "Tingkat Risiko", "Referensi ID", "Tanggal", "Selisih Nominal (Rp)", "Penyebab Teridentifikasi", "Rekomendasi Tindakan"]
    ws4.append([])
    ws4.append([""] + headers4)
    
    for c_idx in range(2, 2 + len(headers4)):
        cell = ws4.cell(row=2, column=c_idx)
        cell.font = font_header
        cell.fill = fill_navy
        cell.alignment = align_center
        cell.border = border_thin
        
    for d in discrepancies:
        r_vals = [
            d.get("direction"), d.get("channel"), d.get("issue_type"), d.get("severity"),
            d.get("reference_id"), d.get("date"), d.get("discrepancy_amount"),
            d.get("probable_cause"), d.get("recommended_action")
        ]
        ws4.append([""] + r_vals)
        cur_row = ws4.max_row
        for c_idx in range(2, 2 + len(headers4)):
            cell = ws4.cell(row=cur_row, column=c_idx)
            cell.font = font_regular
            cell.border = border_thin
            if c_idx == 8:
                cell.number_format = '#,##0.00'
                cell.alignment = align_right
            elif c_idx == 5:
                sev = d.get("severity")
                cell.fill = fill_red if sev == "HIGH" else (fill_yellow if sev == "MEDIUM" else fill_gray)
                cell.alignment = align_center

    # =========================================================
    # TAB 5: DRAFT JURNAL PENYESUAIAN
    # =========================================================
    ws5 = wb.create_sheet(title="Draft Jurnal Penyesuaian")
    ws5.views.sheetView[0].showGridLines = True
    
    headers5 = ["Tanggal", "Kode Akun", "Nama Akun", "Debit (Rp)", "Kredit (Rp)", "Keterangan / Memo"]
    ws5.append([])
    ws5.append([""] + headers5)
    
    for c_idx in range(2, 2 + len(headers5)):
        cell = ws5.cell(row=2, column=c_idx)
        cell.font = font_header
        cell.fill = fill_blue
        cell.alignment = align_center
        cell.border = border_thin
        
    # Buat baris jurnal penyesuaian otomatis
    journals = generate_journal_entries_data(result)
    for j in journals:
        ws5.append(["", j["date"], j["account_code"], j["account_name"], j["debit"], j["credit"], j["memo"]])
        cur_row = ws5.max_row
        for c_idx in range(2, 2 + len(headers5)):
            cell = ws5.cell(row=cur_row, column=c_idx)
            cell.font = font_regular
            cell.border = border_thin
            if c_idx in [5, 6]:
                cell.number_format = '#,##0.00'
                cell.alignment = align_right

    # Auto-adjust column width for all sheets
    for ws in [ws1, ws2, ws3, ws4, ws5]:
        for col in ws.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                if cell.value:
                    val_str = str(cell.value)
                    if len(val_str) > max_len:
                        max_len = len(val_str)
            ws.column_dimensions[col_letter].width = min(max(max_len + 3, 12), 45)
            
    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output

def generate_journal_entries_data(result: Dict[str, Any]) -> List[Dict[str, Any]]:
    """Membuat baris jurnal penyesuaian akuntansi otomatis berbasis hasil rekonsiliasi."""
    journals = []
    discrepancies = result.get("discrepancies", [])
    
    for d in discrepancies:
        dt = d.get("date", "2026-08-15")
        amt = abs(d.get("discrepancy_amount", 0.0))
        issue = d.get("issue_type")
        ref = d.get("reference_id")
        
        if issue == "UNRECORDED_BANK_INTEREST":
            journals.append({"date": dt, "account_code": "1-10001", "account_name": "Kas di Bank BCA", "debit": amt, "credit": 0.0, "memo": f"Penerimaan Bunga Bank ({ref})"})
            journals.append({"date": dt, "account_code": "8-10001", "account_name": "Pendapatan Bunga Giro", "debit": 0.0, "credit": amt, "memo": f"Penerimaan Bunga Bank ({ref})"})
            
        elif issue == "UNRECORDED_BANK_FEE":
            journals.append({"date": dt, "account_code": "6-20001", "account_name": "Beban Administrasi Bank", "debit": amt, "credit": 0.0, "memo": f"Biaya Adm Rekening Koran ({ref})"})
            journals.append({"date": dt, "account_code": "1-10001", "account_name": "Kas di Bank BCA", "debit": 0.0, "credit": amt, "memo": f"Biaya Adm Rekening Koran ({ref})"})
            
        elif issue == "UNRECORDED_BANK_TAX":
            journals.append({"date": dt, "account_code": "6-20002", "account_name": "Beban Pajak Bunga Bank", "debit": amt, "credit": 0.0, "memo": f"Pajak Bunga Giro 20% ({ref})"})
            journals.append({"date": dt, "account_code": "1-10001", "account_name": "Kas di Bank BCA", "debit": 0.0, "credit": amt, "memo": f"Pajak Bunga Giro 20% ({ref})"})
            
        elif issue == "UNIDENTIFIED_BANK_CREDIT":
            journals.append({"date": dt, "account_code": "1-10001", "account_name": "Kas di Bank BCA", "debit": amt, "credit": 0.0, "memo": f"Dana Masuk Belum Teridentifikasi ({ref})"})
            journals.append({"date": dt, "account_code": "2-90001", "account_name": "Titipan / Dana Belum Teridentifikasi", "debit": 0.0, "credit": amt, "memo": f"Dana Masuk Belum Teridentifikasi ({ref})"})
            
        elif issue == "CASH_SHORTAGE_ANOMALY":
            journals.append({"date": dt, "account_code": "6-90001", "account_name": "Beban Selisih Kas (Shortage)", "debit": amt, "credit": 0.0, "memo": f"Selisih Kas Fisik Setoran CDM ({ref})"})
            journals.append({"date": dt, "account_code": "1-10002", "account_name": "Kas Tunai Toko", "debit": 0.0, "credit": amt, "memo": f"Selisih Kas Fisik Setoran CDM ({ref})"})
            
    return journals

def generate_journal_csv(result: Dict[str, Any]) -> str:
    """Mengonversi baris jurnal ke format CSV siap impor Mekari Jurnal / Accurate."""
    output = io.StringIO()
    writer = csv.writer(output)
    writer.writerow(["Date", "Account_Code", "Account_Name", "Debit", "Credit", "Description"])
    
    journals = generate_journal_entries_data(result)
    for j in journals:
        writer.writerow([j["date"], j["account_code"], j["account_name"], j["debit"], j["credit"], j["memo"]])
        
    return output.getvalue()
